#!/usr/bin/env python3
"""Import Toronto population denominators to Markdown.

Produces three files used by the v1 build to compute rates per 10,000 residents
(see ``project/decisions/002-present-rates-not-counts.md``):

1. ``data/population/toronto-neighbourhoods-2021.md`` — one row per neighbourhood
   from the City of Toronto Neighbourhood Profiles 2021 (158-model) XLSX. The
   canonical population column is the row labelled
   ``Total - Age groups of the population - 25% sample data`` per ADR-007 —
   the 158-model release does NOT carry a 100% Census population row.
2. ``data/population/toronto-csd-annual.md`` — one row per reference year
   2016–2025 from StatCan table 17-10-0155-01, filtered to the City of Toronto
   Census Subdivision (DGUID ``2021A00053520005``) per ADR-006.
3. ``data/population/validation.json`` — mirrors the shape of
   ``data/crimes/tps-csi-validation.json``: source URLs, fetch timestamps, row
   counts, spot-checks, and the contract checks described in
   ``docs/data-sources.md`` § Population denominators.

Validation contract (per the revised contract in ``docs/data-sources.md`` and
ADR-007):

- Neighbourhood file: exactly 158 rows; every ``neighbourhood_id`` in the
  imported CSI corpus must appear; sum within ±2% of the City of Toronto 100%
  2021 Census total (2,794,356). The ±2% tolerance accommodates the documented
  ~1.18% under-count of the 25% long-form sample.
- CSD annual file: exactly 10 rows (2016–2025); ``year`` is unique; every
  ``population_july1`` value within [2,500,000, 3,500,000]. Strict YoY
  monotonicity is NOT enforced — the documented 2020→2021 post-Census
  rebasing dip is accepted.

Exits non-zero with a specific error if any check fails.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
import tempfile
import urllib.request
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl

from _md_table import iter_markdown_rows


NEIGHBOURHOOD_XLSX_URL = (
    "https://ckan0.cf.opendata.inter.prod-toronto.ca/dataset/"
    "6e19a90f-971c-46b3-852c-0c48c436d1fc/resource/"
    "19d4a806-7385-4889-acf2-256f1e079060/download/"
    "nbhd_2021_census_profile_full_158model.xlsx"
)

CSD_ZIP_URL = "https://www150.statcan.gc.ca/n1/tbl/csv/17100155-eng.zip"

TORONTO_CSD_DGUID = "2021A00053520005"
TORONTO_2021_CENSUS_TOTAL = 2_794_356
NEIGHBOURHOOD_SUM_TOLERANCE = 0.02  # ±2% per ADR-007

CSD_YEARS = tuple(range(2016, 2026))
CSD_POPULATION_RANGE = (2_500_000, 3_500_000)

NEIGHBOURHOOD_POPULATION_ROW_LABEL = (
    "Total - Age groups of the population - 25% sample data"
)
NEIGHBOURHOOD_NAME_ROW_LABEL = "Neighbourhood Name"
NEIGHBOURHOOD_NUMBER_ROW_LABEL = "Neighbourhood Number"

NEIGHBOURHOOD_COLUMNS = (
    "neighbourhood_id",
    "neighbourhood_name",
    "population_2021_sample25",
)
CSD_COLUMNS = (
    "year",
    "population_july1",
    "source_release_date",
)

USER_AGENT = "crime-pulse-import/1.0"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Import Toronto population denominators to Markdown."
    )
    parser.add_argument(
        "--neighbourhood-url",
        default=NEIGHBOURHOOD_XLSX_URL,
        help="URL of the 158-model neighbourhood profiles XLSX.",
    )
    parser.add_argument(
        "--csd-url",
        default=CSD_ZIP_URL,
        help="URL of the StatCan table 17-10-0155-01 ZIP archive.",
    )
    parser.add_argument(
        "--out-dir",
        default="data/population",
        type=Path,
        help="Directory for generated Markdown and validation files.",
    )
    parser.add_argument(
        "--crimes-dir",
        default="data/crimes",
        type=Path,
        help="Directory of imported CSI Markdown tables (for the join check).",
    )
    return parser.parse_args()


# ---------- fetch helpers ----------


def http_get(url: str, dest: Path) -> None:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=120) as response, dest.open("wb") as handle:
        chunk = response.read(1 << 16)
        while chunk:
            handle.write(chunk)
            chunk = response.read(1 << 16)


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None).isoformat() + "Z"


# ---------- neighbourhood file ----------


def find_label_row(ws, label: str) -> int:
    for row_idx in range(1, ws.max_row + 1):
        value = ws.cell(row_idx, 1).value
        if isinstance(value, str) and value.strip() == label:
            return row_idx
    raise RuntimeError(f"Could not find row labelled {label!r} in neighbourhood XLSX.")


def read_neighbourhoods(xlsx_path: Path) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=False)
    if "hd2021_census_profile" not in workbook.sheetnames:
        raise RuntimeError(
            "Expected sheet 'hd2021_census_profile' in neighbourhood XLSX, "
            f"got {workbook.sheetnames!r}."
        )
    ws = workbook["hd2021_census_profile"]

    name_row = find_label_row(ws, NEIGHBOURHOOD_NAME_ROW_LABEL)
    number_row = find_label_row(ws, NEIGHBOURHOOD_NUMBER_ROW_LABEL)
    population_row = find_label_row(ws, NEIGHBOURHOOD_POPULATION_ROW_LABEL)

    rows: list[dict[str, object]] = []
    for col in range(2, ws.max_column + 1):
        nid = ws.cell(number_row, col).value
        name = ws.cell(name_row, col).value
        pop = ws.cell(population_row, col).value
        if nid is None or name is None or pop is None:
            continue
        try:
            nid_int = int(nid)
            pop_int = int(pop)
        except (TypeError, ValueError) as exc:
            raise RuntimeError(
                f"Non-integer neighbourhood ID/population at column {col}: "
                f"id={nid!r}, pop={pop!r}"
            ) from exc
        rows.append(
            {
                "neighbourhood_id": nid_int,
                "neighbourhood_name": re.sub(r"\s+", " ", str(name)).strip(),
                "population_2021_sample25": pop_int,
            }
        )
    rows.sort(key=lambda r: r["neighbourhood_id"])
    return rows


# ---------- CSD file ----------


def read_csd_annual(zip_path: Path) -> tuple[list[dict[str, object]], str]:
    with zipfile.ZipFile(zip_path) as archive:
        data_name = next(
            (n for n in archive.namelist() if n.endswith(".csv") and "MetaData" not in n),
            None,
        )
        if data_name is None:
            raise RuntimeError(
                f"Data CSV not found inside CSD archive {zip_path}; "
                f"members were {archive.namelist()!r}"
            )
        info = archive.getinfo(data_name)
        release_date = date(*info.date_time[:3]).isoformat()
        with archive.open(data_name) as raw:
            text = io.TextIOWrapper(raw, encoding="utf-8-sig", newline="")
            reader = csv.DictReader(text)
            by_year: dict[int, int] = {}
            for source_row in reader:
                if source_row.get("DGUID") != TORONTO_CSD_DGUID:
                    continue
                ref_date = source_row.get("REF_DATE", "").strip()
                if not ref_date.isdigit():
                    continue
                year = int(ref_date)
                if year not in CSD_YEARS:
                    continue
                value = source_row.get("VALUE", "").strip()
                if not value:
                    raise RuntimeError(
                        f"Missing VALUE for Toronto CSD reference year {year}."
                    )
                # Estimates are published as whole persons, but tolerate float.
                pop = int(round(float(value)))
                if year in by_year and by_year[year] != pop:
                    raise RuntimeError(
                        f"Conflicting estimates for {year}: "
                        f"{by_year[year]} vs {pop}."
                    )
                by_year[year] = pop

    rows = [
        {
            "year": year,
            "population_july1": by_year[year],
            "source_release_date": release_date,
        }
        for year in sorted(by_year)
    ]
    return rows, release_date


# ---------- Markdown writers ----------


def markdown_cell(value: object) -> str:
    if value is None:
        return ""
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text.replace("|", "\\|")


def write_markdown_table(
    path: Path,
    title: str,
    source_line: str,
    columns: tuple[str, ...],
    rows: list[dict[str, object]],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write(f"# {title}\n\n")
        handle.write(source_line + "\n\n")
        handle.write("| " + " | ".join(columns) + " |\n")
        handle.write("| " + " | ".join(["---"] * len(columns)) + " |\n")
        for row in rows:
            handle.write(
                "| " + " | ".join(markdown_cell(row[col]) for col in columns) + " |\n"
            )


# ---------- CSI corpus cross-check ----------


_NEIGHBOURHOOD_ID_RE = re.compile(r"\((\d+)\)\s*$")


def collect_csi_neighbourhood_ids(crimes_dir: Path) -> set[int]:
    if not crimes_dir.is_dir():
        raise RuntimeError(
            f"CSI corpus directory {crimes_dir} not found; cannot run join check."
        )
    ids: set[int] = set()
    md_files = sorted(crimes_dir.glob("tps-csi-*.md"))
    if not md_files:
        raise RuntimeError(
            f"No tps-csi-*.md files found under {crimes_dir}; "
            "import the CSI corpus first."
        )
    for path in md_files:
        # Use the shared, escape-aware Markdown parser and address the
        # neighbourhood column by its header name ("neighbourhood") rather
        # than a hardcoded positional index — robust to column reordering and
        # to pipe-bearing cells (e.g. an offence string containing "\|").
        for row in iter_markdown_rows(path):
            cell = row.get("neighbourhood", "").strip()
            match = _NEIGHBOURHOOD_ID_RE.search(cell)
            if match:
                ids.add(int(match.group(1)))
    return ids


# ---------- validation ----------


def validate_neighbourhoods(
    rows: list[dict[str, object]],
    csi_ids: set[int],
) -> dict[str, object]:
    errors: list[str] = []
    if len(rows) != 158:
        errors.append(
            f"Expected 158 neighbourhood rows, got {len(rows)}."
        )
    ids = [row["neighbourhood_id"] for row in rows]
    if len(set(ids)) != len(ids):
        dupes = sorted({i for i in ids if ids.count(i) > 1})
        errors.append(f"Duplicate neighbourhood IDs: {dupes!r}")

    file_ids = set(ids)
    missing_from_population = sorted(csi_ids - file_ids)
    if missing_from_population:
        errors.append(
            "CSI corpus references neighbourhood IDs missing from the population "
            f"file: {missing_from_population!r}"
        )
    extra_in_population = sorted(file_ids - csi_ids)
    # Extras are not a hard failure (a neighbourhood could legitimately have zero
    # CSI incidents in 2016–2025) but they are surprising — record them.

    total = sum(int(row["population_2021_sample25"]) for row in rows)
    delta = total - TORONTO_2021_CENSUS_TOTAL
    pct = delta / TORONTO_2021_CENSUS_TOTAL
    within_tolerance = abs(pct) <= NEIGHBOURHOOD_SUM_TOLERANCE
    if not within_tolerance:
        errors.append(
            f"Sum of 158 populations {total:,} is "
            f"{pct * 100:+.3f}% from Census total {TORONTO_2021_CENSUS_TOTAL:,}; "
            f"tolerance is ±{NEIGHBOURHOOD_SUM_TOLERANCE * 100:.1f}%."
        )

    return {
        "row_count": len(rows),
        "row_count_ok": len(rows) == 158,
        "csi_ids_count": len(csi_ids),
        "csi_ids_missing_from_population": missing_from_population,
        "csi_ids_extra_in_population": extra_in_population,
        "158_sum_vs_census": {
            "actual": total,
            "target": TORONTO_2021_CENSUS_TOTAL,
            "delta": delta,
            "delta_pct": round(pct * 100, 4),
            "tolerance_pct": NEIGHBOURHOOD_SUM_TOLERANCE * 100,
            "within_tolerance": within_tolerance,
        },
        "errors": errors,
    }


def validate_csd(rows: list[dict[str, object]]) -> dict[str, object]:
    errors: list[str] = []
    if len(rows) != len(CSD_YEARS):
        errors.append(
            f"Expected {len(CSD_YEARS)} CSD rows ({min(CSD_YEARS)}–{max(CSD_YEARS)}), "
            f"got {len(rows)}."
        )
    years = [int(row["year"]) for row in rows]
    if len(set(years)) != len(years):
        dupes = sorted({y for y in years if years.count(y) > 1})
        errors.append(f"Duplicate CSD years: {dupes!r}")
    missing_years = sorted(set(CSD_YEARS) - set(years))
    if missing_years:
        errors.append(f"Missing CSD years: {missing_years!r}")
    extra_years = sorted(set(years) - set(CSD_YEARS))
    if extra_years:
        errors.append(f"Unexpected CSD years: {extra_years!r}")

    lo, hi = CSD_POPULATION_RANGE
    passed: list[int] = []
    failed: list[dict[str, object]] = []
    for row in rows:
        pop = int(row["population_july1"])
        if lo <= pop <= hi:
            passed.append(int(row["year"]))
        else:
            failed.append({"year": int(row["year"]), "population_july1": pop})
    if failed:
        errors.append(
            f"CSD population outside range [{lo:,}, {hi:,}] for: "
            + ", ".join(f"{f['year']}={f['population_july1']:,}" for f in failed)
        )

    return {
        "row_count": len(rows),
        "row_count_ok": len(rows) == len(CSD_YEARS),
        "year_range_check": {
            "range": [lo, hi],
            "passed_years": sorted(passed),
            "failed_years": failed,
        },
        "errors": errors,
    }


# ---------- main ----------


def main() -> int:
    args = parse_args()
    args.out_dir.mkdir(parents=True, exist_ok=True)

    neighbourhood_md = args.out_dir / "toronto-neighbourhoods-2021.md"
    csd_md = args.out_dir / "toronto-csd-annual.md"
    validation_json = args.out_dir / "validation.json"

    with tempfile.TemporaryDirectory(prefix="crime-pulse-pop-") as tmpdir:
        tmp_path = Path(tmpdir)

        # 1. Neighbourhood XLSX
        print(f"Fetching neighbourhood XLSX from {args.neighbourhood_url}")
        xlsx_path = tmp_path / "nbhd_2021_158model.xlsx"
        http_get(args.neighbourhood_url, xlsx_path)
        neighbourhood_fetched_at = utc_now_iso()
        print(f"  saved to {xlsx_path} ({xlsx_path.stat().st_size:,} bytes)")

        print("Parsing neighbourhood populations (25% sample data row)")
        neighbourhood_rows = read_neighbourhoods(xlsx_path)
        print(f"  parsed {len(neighbourhood_rows)} neighbourhoods")

        # 2. CSD ZIP
        print(f"Fetching CSD archive from {args.csd_url}")
        zip_path = tmp_path / "17100155-eng.zip"
        http_get(args.csd_url, zip_path)
        csd_fetched_at = utc_now_iso()
        print(f"  saved to {zip_path} ({zip_path.stat().st_size:,} bytes)")

        print(f"Filtering CSD table to DGUID {TORONTO_CSD_DGUID}, "
              f"years {min(CSD_YEARS)}–{max(CSD_YEARS)}")
        csd_rows, csd_release_date = read_csd_annual(zip_path)
        print(f"  parsed {len(csd_rows)} rows; release date {csd_release_date}")

    # 3. CSI cross-check
    print(f"Cross-checking against CSI corpus under {args.crimes_dir}")
    csi_ids = collect_csi_neighbourhood_ids(args.crimes_dir)
    print(f"  {len(csi_ids)} distinct neighbourhood IDs in CSI corpus")

    # 4. Validation
    neighbourhood_validation = validate_neighbourhoods(neighbourhood_rows, csi_ids)
    csd_validation = validate_csd(csd_rows)

    errors = list(neighbourhood_validation["errors"]) + list(csd_validation["errors"])
    if errors:
        print("Validation FAILED; existing population files were not changed:", file=sys.stderr)
        for err in errors:
            print(f"  - {err}", file=sys.stderr)
        return 1

    # 5. Spot-checks for the validation report
    def find_neighbourhood(rows: list[dict[str, object]], nid: int) -> dict[str, object] | None:
        for row in rows:
            if row["neighbourhood_id"] == nid:
                return row
        return None

    def find_year(rows: list[dict[str, object]], year: int) -> dict[str, object] | None:
        for row in rows:
            if row["year"] == year:
                return row
        return None

    moss_park = find_neighbourhood(neighbourhood_rows, 73)
    annex = find_neighbourhood(neighbourhood_rows, 95)
    toronto_csd_2021 = find_year(csd_rows, 2021)
    toronto_csd_2025 = find_year(csd_rows, 2025)

    report = {
        "neighbourhood_file": {
            "source": {
                "name": (
                    "City of Toronto Neighbourhood Profiles 2021 (158-model)"
                ),
                "url": args.neighbourhood_url,
                "fetched_at": neighbourhood_fetched_at,
                "population_row_label": NEIGHBOURHOOD_POPULATION_ROW_LABEL,
            },
            "output_path": str(neighbourhood_md),
            "row_count": len(neighbourhood_rows),
            "validation": neighbourhood_validation,
            "spot_checks": {
                "moss_park_73": (
                    None if moss_park is None else moss_park["population_2021_sample25"]
                ),
                "annex_95": (
                    None if annex is None else annex["population_2021_sample25"]
                ),
            },
        },
        "csd_file": {
            "source": {
                "name": (
                    "Statistics Canada Table 17-10-0155-01 — "
                    "Population estimates, July 1, by census subdivision, "
                    "2021 boundaries"
                ),
                "url": args.csd_url,
                "fetched_at": csd_fetched_at,
                "filter_dguid": TORONTO_CSD_DGUID,
                "release_date": csd_release_date,
            },
            "output_path": str(csd_md),
            "row_count": len(csd_rows),
            "validation": csd_validation,
            "spot_checks": {
                "toronto_csd_2021": (
                    None
                    if toronto_csd_2021 is None
                    else toronto_csd_2021["population_july1"]
                ),
                "toronto_csd_2025": (
                    None
                    if toronto_csd_2025 is None
                    else toronto_csd_2025["population_july1"]
                ),
            },
        },
    }
    # 6. Stage all three outputs before replacing any canonical file. This
    # keeps validation and preparation failures non-destructive and gives each
    # final replacement same-filesystem atomicity.
    with tempfile.TemporaryDirectory(
        prefix=".crime-pulse-pop-output-", dir=args.out_dir
    ) as output_tmp:
        stage = Path(output_tmp)
        staged_neighbourhood = stage / neighbourhood_md.name
        staged_csd = stage / csd_md.name
        staged_validation = stage / validation_json.name

        write_markdown_table(
            staged_neighbourhood,
            title="Toronto Neighbourhood Populations 2021 (158-model, 25% sample)",
            source_line=(
                "Source: City of Toronto Neighbourhood Profiles 2021 (158-model), "
                f"`{args.neighbourhood_url}`. Population column is the row "
                f"`{NEIGHBOURHOOD_POPULATION_ROW_LABEL}`; "
                "the 158-model XLSX does not publish a 100% Census population row "
                "(see `project/decisions/007-per-neighbourhood-denominator-is-25pct-sample.md`)."
            ),
            columns=NEIGHBOURHOOD_COLUMNS,
            rows=neighbourhood_rows,
        )
        write_markdown_table(
            staged_csd,
            title="Toronto Census Subdivision Population, July 1, 2016–2025",
            source_line=(
                "Source: Statistics Canada Table 17-10-0155-01, "
                f"`{args.csd_url}`. Filtered to City of Toronto CSD "
                f"(DGUID `{TORONTO_CSD_DGUID}`); see "
                "`project/decisions/006-city-wide-denominator-is-toronto-csd-not-cma.md`."
            ),
            columns=CSD_COLUMNS,
            rows=csd_rows,
        )
        staged_validation.write_text(
            json.dumps(report, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )

        staged_neighbourhood.replace(neighbourhood_md)
        staged_csd.replace(csd_md)
        staged_validation.replace(validation_json)

    sum_check = neighbourhood_validation["158_sum_vs_census"]
    print()
    print("Population import complete.")
    print(f"  Neighbourhood file: {neighbourhood_md} ({len(neighbourhood_rows)} rows)")
    print(
        "  Neighbourhood 158-sum: "
        f"{sum_check['actual']:,} vs Census {sum_check['target']:,} "
        f"({sum_check['delta_pct']:+.3f}%)"
    )
    print(f"  CSD file: {csd_md} ({len(csd_rows)} rows)")
    if toronto_csd_2025 is not None:
        print(
            f"  Toronto CSD 2025 estimate: "
            f"{toronto_csd_2025['population_july1']:,} (July 1)"
        )
    print(f"  Validation report: {validation_json}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
